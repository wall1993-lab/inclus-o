"""
Corretor - Prepara ENEM
Peca 3 de 3 do sistema (gerador -> leitor -> corretor).

Recebe o gabarito (a resposta certa de cada questao) e as respostas que
a Parte 2 (leitor) leu de cada aluno, compara as duas coisas, calcula a
nota, e gera uma planilha (.xlsx) com a lista de alunos e as notas.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def parse_gabarito(texto, num_questoes):
    """Converte o texto digitado/colado pela coordenadora num dicionario
    {"1": "A", "2": "B", ...}. Aceita 3 formatos:

      1) uma letra por linha, na ordem das questoes (1, 2, 3...)
      2) "numero: letra" ou "numero, letra" por linha (pode pular questao)
      3) tudo numa linha so', sem separador: "ABCDE..." (uma letra por questao)

    Devolve (gabarito, erros).
    """
    texto = texto.strip()
    erros = []
    gabarito = {}

    if not texto:
        return {}, ["O gabarito esta vazio."]

    linhas = [l.strip() for l in texto.splitlines() if l.strip()]

    # formato 3: tudo numa linha so', do tamanho do numero de questoes
    if len(linhas) == 1 and (":" not in linhas[0]) and ("," not in linhas[0]) and " " not in linhas[0]:
        letras = linhas[0].upper()
        if len(letras) == num_questoes:
            for i, letra in enumerate(letras, start=1):
                gabarito[str(i)] = letra
            return gabarito, []

    # formato 2: "numero: letra" ou "numero, letra"
    if any(":" in l or "," in l for l in linhas):
        for numero_linha, linha in enumerate(linhas, start=1):
            separador = ":" if ":" in linha else ","
            partes = [p.strip() for p in linha.split(separador)]
            if len(partes) != 2 or not partes[0].isdigit():
                erros.append(f"Linha {numero_linha}: nao entendi \"{linha}\" (esperava \"numero: letra\").")
                continue
            gabarito[partes[0]] = partes[1].upper()
        return gabarito, erros

    # formato 1: uma letra por linha
    for i, linha in enumerate(linhas, start=1):
        gabarito[str(i)] = linha.upper()

    if len(gabarito) != num_questoes:
        erros.append(
            f"O gabarito tem {len(gabarito)} questao(oes), mas o simulado tem {num_questoes}."
        )

    return gabarito, erros


def corrigir(alunos, respostas_por_aluno, gabarito, num_questoes):
    """Compara as respostas de cada aluno com o gabarito.

    'alunos' e' a lista [{"nome", "serie", "matricula", "turma"?}, ...]
    'respostas_por_aluno' vem da Parte 2 (leitor): {matricula: {"1": "A", ...}}
    'gabarito' e' {"1": "A", "2": "B", ...}

    Devolve uma lista de dicionarios, um por aluno, com nome/serie/matricula,
    acertos, questoes em branco, questoes com dupla marcacao, nota (0 a 10)
    e a resposta que o aluno deu em cada questao.
    """
    resultado = []
    for aluno in alunos:
        matricula = aluno["matricula"]
        respostas = respostas_por_aluno.get(matricula)

        linha = {
            "nome": aluno["nome"],
            "serie": aluno.get("serie", ""),
            "turma": aluno.get("turma", ""),
            "matricula": matricula,
            "recebeu_leitura": respostas is not None,
        }

        acertos = 0
        em_branco = 0
        duplas = 0
        respostas_por_questao = {}

        for questao in range(1, num_questoes + 1):
            chave = str(questao)
            certa = gabarito.get(chave)
            dada = (respostas or {}).get(chave)
            respostas_por_questao[chave] = dada

            if dada is None:
                em_branco += 1
            elif dada == "MULTIPLA":
                duplas += 1
            elif certa is not None and dada == certa:
                acertos += 1

        linha["acertos"] = acertos
        linha["em_branco"] = em_branco
        linha["duplas_marcacoes"] = duplas
        linha["total_questoes"] = num_questoes
        linha["nota"] = round((acertos / num_questoes) * 10, 2) if num_questoes else 0
        linha["respostas_por_questao"] = respostas_por_questao

        resultado.append(linha)

    return resultado


def gerar_planilha(resultado_corrigido, num_questoes, saida_xlsx="notas.xlsx"):
    """Gera o arquivo .xlsx com a lista de alunos, notas e o detalhe de
    resposta questao a questao."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notas"

    colunas_fixas = [
        "Nome", "Serie", "Turma", "Matricula", "Acertos", "Total questoes",
        "Nota", "Em branco", "Dupla marcacao", "Cartao lido?",
    ]
    colunas_questoes = [f"Q{n}" for n in range(1, num_questoes + 1)]
    cabecalho = colunas_fixas + colunas_questoes

    ws.append(cabecalho)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")
        celula.alignment = Alignment(horizontal="center")

    for linha in resultado_corrigido:
        valores = [
            linha["nome"], linha["serie"], linha["turma"], linha["matricula"],
            linha["acertos"], linha["total_questoes"], linha["nota"],
            linha["em_branco"], linha["duplas_marcacoes"],
            "Sim" if linha["recebeu_leitura"] else "NAO",
        ]
        for n in range(1, num_questoes + 1):
            dada = linha["respostas_por_questao"].get(str(n))
            valores.append(dada if dada else ("-" if dada is None else dada))
        ws.append(valores)

    ws.freeze_panes = "E2"
    for i, titulo in enumerate(cabecalho, start=1):
        largura = max(10, len(titulo) + 2) if i <= len(colunas_fixas) else 5
        ws.column_dimensions[get_column_letter(i)].width = largura

    wb.save(saida_xlsx)
    return saida_xlsx
